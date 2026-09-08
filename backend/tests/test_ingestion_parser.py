import io
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
from PIL import Image

from app.ingestion.chunker import chunk_pages
from app.ingestion.parser import parse_docx


def test_docx_tables_keep_their_original_case_order(tmp_path):
    doc = Document()
    for number, marker in [(1, "FIRST_CASE_ALT_42"), (2, "SECOND_CASE_ALT_73")]:
        doc.add_paragraph(f"Case {number}")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = marker
    path = tmp_path / "cases.docx"
    doc.save(path)

    pages = parse_docx(str(path))
    chunks = chunk_pages(pages)

    assert pages[0].text.index("FIRST_CASE_ALT_42") < pages[0].text.index("Case 2")
    first = next(chunk.content for chunk in chunks if "Case 1" in chunk.content)
    second = next(chunk.content for chunk in chunks if "Case 2" in chunk.content)
    assert "FIRST_CASE_ALT_42" in first
    assert "SECOND_CASE_ALT_73" not in first
    assert "SECOND_CASE_ALT_73" in second
    assert "FIRST_CASE_ALT_42" not in second


def test_wps_repair_preserves_text_and_embedded_images(tmp_path):
    picture = io.BytesIO()
    Image.new("RGB", (8, 8), "blue").save(picture, format="PNG")
    picture.seek(0)
    doc = Document()
    doc.add_paragraph("合成病例正文")
    doc.add_picture(picture)
    original = io.BytesIO()
    doc.save(original)
    path = tmp_path / "wps.docx"
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                root = ET.fromstring(data)
                ET.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
                    "Id": "rIdBroken",
                    "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                    "Target": "NULL",
                })
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
    original_bytes = path.read_bytes()

    pages = parse_docx(str(path))

    assert pages[0].text == "合成病例正文"
    assert len(pages[0].images) == 1
    assert pages[0].images[0].blob == picture.getvalue()
    assert path.read_bytes() == original_bytes
